"""Generate persistent Word delivery reports for completed deployments."""
from __future__ import annotations

import os
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.config import settings


INK = "0B2545"
BLUE = "2E74B5"
MUTED = "64748B"
LIGHT = "E8EEF5"
PALE = "F4F6F9"
CAUTION = "FFF4CE"
CAUTION_TEXT = "7A5A00"
WHITE = "FFFFFF"
GRID = "CBD5E1"
FONT_ASCII = "Calibri"
FONT_CJK = "Microsoft YaHei"
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def _set_run_font(
    run,
    *,
    size: float = 11,
    color: str = INK,
    bold: bool = False,
    italic: bool = False,
) -> None:
    run.font.name = FONT_ASCII
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), FONT_ASCII)
    fonts.set(qn("w:hAnsi"), FONT_ASCII)
    fonts.set(qn("w:eastAsia"), FONT_CJK)


def _set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def _set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths_dxa: list[int]) -> None:
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError("Word table column widths must sum to 9360 DXA")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[index] / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_margins(cell)


def _set_table_borders(table, color: str = GRID, size: int = 4) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def _format_cell(cell, text: str, *, bold=False, color=INK, size=9.5, align=None) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.1
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(str(text))
    _set_run_font(run, size=size, color=color, bold=bold)


def _add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.add_run(text)


def _add_key_value_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=0, cols=2)
    for label, value in rows:
        cells = table.add_row().cells
        _set_cell_fill(cells[0], PALE)
        _format_cell(cells[0], label, bold=True, size=9.5)
        _format_cell(cells[1], value, size=9.5)
    _set_table_geometry(table, [2700, 6660])
    _set_table_borders(table)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def _add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    _set_run_font(run, size=9, color=MUTED)


def _configure_document(document: Document, task_id: str) -> None:
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = FONT_ASCII
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for style_name, size, before, after, color in (
        ("Heading 1", 16, 16, 8, BLUE),
        ("Heading 2", 13, 12, 6, BLUE),
        ("Heading 3", 12, 8, 4, INK),
    ):
        style = styles[style_name]
        style.font.name = FONT_ASCII
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _set_run_font(header.add_run("SP MediDeploy  |  Redis 部署交付报告"), size=9, color=MUTED)
    header.add_run(" " * 8)
    _set_run_font(header.add_run(task_id), size=8, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_run_font(footer.add_run("SPMP · 第 "), size=9, color=MUTED)
    _add_page_field(footer)
    _set_run_font(footer.add_run(" 页"), size=9, color=MUTED)


def build_report(
    snapshot: dict,
    *,
    redis_password: str,
    output_path: Path,
    generated_at: datetime | None = None,
) -> Path:
    generated_at = generated_at or datetime.now()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    _configure_document(document, snapshot["task_id"])
    document.core_properties.title = f"{snapshot['task_name']} - Redis 部署交付报告"
    document.core_properties.subject = "SP MediDeploy deployment delivery report"
    document.core_properties.author = "SP MediDeploy Platform"
    document.core_properties.keywords = "SPMP, Redis, deployment, delivery"

    kicker = document.add_paragraph()
    kicker.paragraph_format.space_after = Pt(3)
    _set_run_font(kicker.add_run("SP MEDIDEPLOY · DELIVERY REPORT"), size=9, color=BLUE, bold=True)
    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    _set_run_font(title.add_run("Redis 部署交付报告"), size=23, color=INK, bold=True)
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    _set_run_font(
        subtitle.add_run(snapshot["task_name"]),
        size=13,
        color=MUTED,
        bold=True,
    )

    _add_key_value_table(
        document,
        [
            ("任务 ID", snapshot["task_id"]),
            ("部署模式", "Redis Cluster" if snapshot["mode"] == "cluster" else "单机实例"),
            ("软件版本", snapshot["package"]["version"]),
            ("软件包", f"{snapshot['package']['package_type']} · {snapshot['package']['filename']}"),
            ("部署完成时间", snapshot["completed_at"]),
            ("报告生成时间", generated_at.strftime("%Y-%m-%d %H:%M:%S")),
        ],
    )

    warning = document.add_table(rows=1, cols=1)
    _set_cell_fill(warning.cell(0, 0), CAUTION)
    _format_cell(
        warning.cell(0, 0),
        "敏感信息提示：本报告包含 Redis 明文访问密码。请仅在授权范围内存储、传输和使用。",
        bold=True,
        color=CAUTION_TEXT,
        size=10,
    )
    _set_table_geometry(warning, [CONTENT_WIDTH_DXA])
    _set_table_borders(warning, color="E5C365", size=6)

    _add_heading(document, "1. 部署概览")
    config = snapshot["config"]
    _add_key_value_table(
        document,
        [
            ("实例数量", str(len(snapshot["instances"]))),
            ("主节点配置数", str(config["primary_count"])),
            ("每主节点副本数", str(config["replicas_per_primary"])),
            ("AOF 持久化", "已开启" if config["appendonly"] else "未开启"),
            ("最大内存", config.get("maxmemory") or "未限制"),
            ("内存淘汰策略", config["maxmemory_policy"]),
        ],
    )

    _add_heading(document, "2. 访问凭据")
    _add_key_value_table(
        document,
        [
            ("Redis 用户", "default"),
            ("Redis 密码", redis_password),
            ("连接示例", f"redis-cli -h {snapshot['instances'][0]['address']} -p {snapshot['instances'][0]['redis_port']} -a <报告中的密码>"),
        ],
    )

    _add_heading(document, "3. 节点与端口")
    topology = document.add_table(rows=1, cols=6)
    headers = ["序号", "服务器", "SSH", "Redis", "Cluster Bus", "任务服务名"]
    for cell, value in zip(topology.rows[0].cells, headers):
        _set_cell_fill(cell, LIGHT)
        _format_cell(cell, value, bold=True, color=INK, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    for index, instance in enumerate(snapshot["instances"], 1):
        cells = topology.add_row().cells
        values = [
            str(index),
            f"{instance['host_name']}\n{instance['address']}",
            str(instance["ssh_port"]),
            str(instance["redis_port"]),
            str(instance["bus_port"]) if snapshot["mode"] == "cluster" else "-",
            instance["service_name"],
        ]
        for column, (cell, value) in enumerate(zip(cells, values)):
            _format_cell(
                cell,
                value,
                size=8.2,
                align=WD_ALIGN_PARAGRAPH.CENTER if column in (0, 2, 3, 4) else WD_ALIGN_PARAGRAPH.LEFT,
            )
    _set_table_geometry(topology, [650, 2250, 850, 850, 1150, 3610])
    _set_table_borders(topology)

    _add_heading(document, "4. 部署目录")
    directories = document.add_table(rows=1, cols=5)
    for cell, value in zip(directories.rows[0].cells, ["节点", "安装目录", "数据目录", "日志目录", "配置目录"]):
        _set_cell_fill(cell, LIGHT)
        _format_cell(cell, value, bold=True, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    for index, instance in enumerate(snapshot["instances"], 1):
        cells = directories.add_row().cells
        for column, value in enumerate(
            [
                str(index),
                instance["install_dir"],
                instance["data_dir"],
                instance["log_dir"],
                instance["config_dir"],
            ]
        ):
            _format_cell(
                cells[column],
                value,
                size=7.8,
                align=WD_ALIGN_PARAGRAPH.CENTER if column == 0 else WD_ALIGN_PARAGRAPH.LEFT,
            )
    _set_table_geometry(directories, [600, 2190, 2190, 2190, 2190])
    _set_table_borders(directories)

    _add_heading(document, "5. 运维说明")
    notes = [
        ("配置快照", "本报告记录的是部署完成时的平台配置快照，后续资产名称或凭据变更不会改写本报告。"),
        ("网络要求", "Redis Cluster 对外服务需同时保证 Redis 端口及对应 Cluster Bus 端口可达。"),
        ("服务管理", "服务由 SPMP 创建的独立 systemd 单元管理；执行变更前请先确认任务 ID 与服务名。"),
        ("密码轮换", "如需轮换密码，应同步更新 Redis 配置并重新生成新的交付记录。"),
    ]
    for label, text in notes:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(5)
        _set_run_font(paragraph.add_run(f"{label}："), size=10, color=BLUE, bold=True)
        _set_run_font(paragraph.add_run(text), size=10, color=INK)

    temporary_path = output_path.with_suffix(".tmp.docx")
    document.save(temporary_path)
    os.replace(temporary_path, output_path)
    return output_path


def report_path_for(deployment_id: str) -> Path:
    return Path(settings().reports_dir) / f"redis-deployment-{deployment_id}.docx"
